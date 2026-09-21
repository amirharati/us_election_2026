# /// script
# requires-python = ">=3.11"
# dependencies = ["xlrd==2.0.2", "openpyxl==3.1.5", "pypdf==6.19.0"]
# ///
"""Offline, source-separated feature preparation. No poll joins or forecast features.

Run with uv run --script scripts/prepare_features.py. Inputs are pinned on first
run; --refresh-inputs explicitly adopts newly downloaded snapshots.
"""
from __future__ import annotations
import argparse
import calendar
from collections import Counter
import csv
from datetime import date, datetime, timedelta
import hashlib
from html.parser import HTMLParser
import io
import json
from pathlib import Path
import re
import sqlite3
import zipfile

from data_utils import LAB, Snapshot, atomic_json, json_bytes

VERSION = 1
BASE = LAB / "data/features"
OUTPUT = LAB / "data/features_prepared"
CONFIG = LAB / "config/feature_preparation_inputs.json"
MISSING = {"", "na", "n/a", "nan", "null", "none", ".", "..", "...", "-"}
SERIES = {
    "CPIAUCSL": ("month", "index_1982_84_100_SA"),
    "CUSR0000SETB01": ("month", "index_1982_84_100_SA"),
    "CUSR0000SA0": ("month", "index_1982_84_100_SA"),
    "UNRATE": ("month", "percent"), "LNS14000000": ("month", "percent"),
    "A229RX0": ("month", "chained_2017_USD_per_capita_SAAR"),
    "GDPC1": ("quarter", "billions_chained_2017_USD_SAAR"),
    "PCEPI": ("month", "index_2017_100_SA"),
    "NASDAQCOM": ("day", "index_1971_02_05_100"),
    "STLFSI4": ("week", "standardized_index"),
    "USACSCICP02STSAM": ("month", "percentage_balance_SA"),
}


def sha_file(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def stable(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False,
                                     separators=(",", ":"), default=str).encode()).hexdigest()


def slug(value):
    return re.sub(r"[^a-z0-9]+", "_", str(value).lower()).strip("_")


def iso(value):
    """Strict dates; do not silently repair malformed dates or guess day/month."""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip()
    if not text:
        return ""
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%Y%m%d", "%b %d, %Y", "%B %d, %Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}[ T]\d\d:\d\d:\d\d(?:\.\d+)?(?:Z)?", text):
        return date.fromisoformat(text[:10]).isoformat()
    raise ValueError(f"Unrecognized date: {text!r}")


def period(value, frequency="day"):
    d = date.fromisoformat(iso(value))
    if frequency == "month":
        start, end = d.replace(day=1), d.replace(day=calendar.monthrange(d.year, d.month)[1])
    elif frequency == "quarter":
        month = 1 + 3 * ((d.month - 1) // 3)
        start = date(d.year, month, 1)
        end = date(d.year, month + 2, calendar.monthrange(d.year, month + 2)[1])
    elif frequency == "year":
        start, end = date(d.year, 1, 1), date(d.year, 12, 31)
    elif frequency == "week":
        start, end = d - timedelta(days=6), d
    elif frequency in ("day", "report_date"):
        start = end = d
    else:
        raise ValueError(frequency)
    return dict(date=d.isoformat(), period_start=start.isoformat(), period_end=end.isoformat(),
                date_precision=frequency, available_at="")


def number(raw, unit="source_numeric", sentinels=()):
    """Return value, normalized unit, status. Never coerce missing/annotations to 0."""
    import math
    text = str(raw).strip() if raw is not None else ""
    outunit = "fraction" if unit == "percent" else unit
    if text.lower() in MISSING:
        return "", outunit, "missing"
    if text == "*" or text.startswith(("<", ">")):
        return "", outunit, "censored"
    try:
        value = float(text.rstrip("%").replace(",", "").replace("−", "-"))
    except ValueError:
        return "", outunit, "unparsed_numeric"
    if not math.isfinite(value):
        return "", outunit, "nonfinite"
    if value in sentinels:
        return "", outunit, "source_missing_code"
    if unit == "percent":
        value /= 100
    return value, outunit, "observed"


class Tables(HTMLParser):
    def __init__(self):
        super().__init__(); self.tables=[]; self.depth=0; self.row=None; self.cell=None
    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self.depth += 1
            if self.depth == 1: self.tables.append([])
        elif self.depth == 1 and tag == "tr": self.row=[]
        elif self.depth == 1 and tag in ("td", "th"): self.cell=[]
        elif self.cell is not None and tag == "br": self.cell.append(" ")
    def handle_data(self, data):
        if self.cell is not None: self.cell.append(data)
    def handle_endtag(self, tag):
        if tag in ("td", "th") and self.cell is not None:
            if self.row is not None: self.row.append(" ".join("".join(self.cell).split()))
            self.cell=None
        elif tag == "tr" and self.depth == 1 and self.row is not None:
            self.tables[-1].append(self.row); self.row=None
        elif tag == "table": self.depth -= 1


class Builder:
    def __init__(self, snap):
        self.snap=snap; self.outputs={}; self.catalog={}; self.counts=Counter(); self.issues=[]
        self.db=sqlite3.connect(snap.stage / "_dedup.sqlite")
        self.db.execute("CREATE TABLE records (table_name TEXT, id TEXT PRIMARY KEY, logical_key TEXT)")
        self.db.execute("CREATE INDEX logical_keys ON records(table_name,logical_key)")
        self.artifacts=[]; self.a=None; self.source=""

    def write(self, name, row, key=None):
        flags = []
        # Flag source inconsistencies without changing or winsorizing observations.
        if self.source in ('michigan','ucsb','policy_agendas') and row.get('unit')=='fraction':
            if row.get('value','')!='' and not 0 <= row['value'] <= 1:
                flags.append('share_out_of_range')
        if self.source=='silver_bulletin':
            for k in ('approve_fraction','disapprove_fraction'):
                if row.get(k,'')!='' and not 0 <= row[k] <= 1: flags.append(k+'_out_of_range')
            a,d,n=(row.get(k,'') for k in ('approve_fraction','disapprove_fraction','net_fraction'))
            if '' not in (a,d,n) and abs(a-d-n)>0.00001: flags.append('net_arithmetic_mismatch')
        if self.source=='ucdp':
            low,best,high=(row.get(k,'') for k in ('low','best','high'))
            if '' not in (low,best,high) and not 0 <= low <= best <= high: flags.append('fatality_bounds_inconsistent')
            if row.get('latitude','')!='' and not -90 <= row['latitude'] <=90: flags.append('latitude_out_of_range')
            if row.get('longitude','')!='' and not -180 <= row['longitude'] <=180: flags.append('longitude_out_of_range')
        if self.source=='cboe':
            low,high,op,close=(row.get(k,'') for k in ('low','high','open','close'))
            if '' not in (low,high,op,close) and not low <= min(op,close) <= max(op,close) <= high:
                flags.append('ohlc_inconsistent')
        if self.source in ('who','owid') and any(isinstance(v,(float,int)) and v<0 for v in row.values()):
            flags.append('negative_value_or_reporting_correction')
        row={**row,'quality_flags':'|'.join(flags)}
        for flag in flags:self.issue(flag,dict(table=name,key=key))
        name=f"{self.source}/{name}"
        fields=list(row)
        if name not in self.outputs:
            p=self.snap.stage/(name+".csv");p.parent.mkdir(parents=True, exist_ok=True)
            f=p.open("w",newline="",encoding="utf-8")
            writer=csv.DictWriter(f,fieldnames=["record_id","logical_key","artifact_id","source_row",*fields],lineterminator="\n")
            writer.writeheader();self.outputs[name]=(f,writer,fields)
            self.catalog[name]={"path":name+".csv","columns":fields,"rows":0,"first_date":"","last_date":""}
        f,w,expected=self.outputs[name]
        if fields != expected: raise ValueError(f"Schema drift in {name}: {fields} vs {expected}")
        rid=stable([name,row]); logical=stable([name, key if key is not None else row])
        found=self.db.execute("SELECT id FROM records WHERE id=?",(rid,)).fetchone()
        if found:
            self.counts[f"{self.source}:exact_duplicate_rows"]+=1
            self.audit("duplicate_rows",dict(canonical_record_id=rid,artifact_id=self.a["id"],source_row=self.line,table=name))
            return
        self.db.execute("INSERT INTO records VALUES (?,?,?)",(name,rid,logical))
        w.writerow(dict(record_id=rid,logical_key=logical,artifact_id=self.a["id"],source_row=self.line,**row))
        c=self.catalog[name];c["rows"]+=1
        d=row.get("date",row.get("period_end",row.get("end_date",row.get("date_end",""))))
        if d:
            c["first_date"]=min(c["first_date"] or d,d);c["last_date"]=max(c["last_date"],d)
        self.counts[f"{self.source}:rows"]+=1

    def audit(self, name, row):
        p=self.snap.stage/"audit"/(name+".jsonl");p.parent.mkdir(exist_ok=True)
        with p.open("a",encoding="utf-8") as f: f.write(json.dumps(row,ensure_ascii=False,default=str)+"\n")

    def issue(self, reason, detail, severity="review"):
        self.counts[f"{self.source}:{reason}"]+=1
        self.audit("issues",dict(source=self.source,artifact_id=self.a["id"],source_row=self.line,
                                 reason=reason,severity=severity,detail=detail))

    def numeric(self, raw, unit="source_numeric", field="", sentinels=()):
        v,u,status=number(raw,unit,sentinels)
        if status not in ("observed", "missing"):
            self.issue(status,dict(field=field,raw=raw))
        if status != "observed": self.counts[f"{self.source}:missing_numeric_cells"]+=1
        return v,u,status

    def observation(self, table, d, frequency, series, raw, unit, group="all", **extra):
        value,u,status=self.numeric(raw,unit,series)
        row={**period(d,frequency),"series":series,"group":group,"value":value,"unit":u,
             "source_value":str(raw),"source_unit":unit,"value_status":status,
             "geography":extra.pop('geography','US'),**extra}
        self.write(table,row,key=[row["date"],series,group,extra])

    def run(self, a):
        self.a=a;self.source=a["source"];self.line=0
        p=LAB/a["path"]
        if sha_file(p) != a["sha256"]: raise ValueError(f"Input checksum mismatch: {p}")
        print(f"Preparing {self.source}/{p.name}",flush=True)
        prepare(self,p)

    def finish(self):
        for name,(f,_,_) in self.outputs.items(): f.close()
        self.db.commit()
        groups=0
        for table,key,n in self.db.execute("SELECT table_name,logical_key,count(*) FROM records GROUP BY table_name,logical_key HAVING count(*)>1"):
            ids=[r[0] for r in self.db.execute("SELECT id FROM records WHERE table_name=? AND logical_key=?",(table,key))]
            self.audit("overlapping_variants",dict(table=table,logical_key=key,variants=n,record_ids=ids));groups+=1
        self.db.close();(self.snap.stage/"_dedup.sqlite").unlink()
        return groups


def rows(path):
    with path.open(encoding="utf-8-sig",newline="") as f:
        rd=csv.DictReader(f)
        if len(rd.fieldnames or []) != len(set(rd.fieldnames or [])): raise ValueError(f"Duplicate headers: {path}")
        for i,r in enumerate(rd,2):
            if None in r or any(v is None for v in r.values()): raise ValueError(f"Malformed CSV: {path}:{i}")
            yield i,r


def workbook(path):
    if path.suffix == ".xls":
        import xlrd
        book=xlrd.open_workbook(path)
        return [(s.name,[s.row_values(i) for i in range(s.nrows)]) for s in book.sheets()]
    import openpyxl
    book=openpyxl.load_workbook(path,read_only=True,data_only=True)
    result=[(s.title,[[v if v is not None else "" for v in row] for row in s.values]) for s in book]
    book.close();return result


def wide(b, table, r, d, freq, numeric, key, drop=(), units=None, sentinels=()):
    """Clean wide native table; dictionary retains source labels and units."""
    units=units or {}; names={k:slug(k) for k in r if k not in drop}
    if len(names.values()) != len(set(names.values())): raise ValueError("Column normalization collision")
    out=period(d,freq); missing={};schema={}
    for old,new in names.items():
        if new in out: new="source_"+new
        raw=r[old];u=units.get(old,"source_numeric")
        if old in numeric:
            v,nu,status=b.numeric(raw,u,old,sentinels);out[new]=v
            # An empty numeric CSV field already means missing. Record exceptional
            # codes separately without repeating hundreds of long field names.
            if status not in ("observed", "missing"): missing[new]=status
            schema[new]={"source_column":old,"source_unit":u,"unit":nu,"type":"number"}
        else:
            out[new]=str(raw).strip();schema[new]={"source_column":old,"type":"string"}
    out["missing_status_json"]=json.dumps(missing,sort_keys=True,separators=(",",":"))
    b.write(table,out,key)
    b.catalog[f"{b.source}/{table}"]["source_columns"]=schema


def prepare(b,p):
    s=b.source
    if s in ("fred","bea","oecd","bls") and p.suffix==".csv":
        for b.line,r in rows(p):
            sid=r.get("series_id",p.stem);freq,u=SERIES[sid]
            b.observation("observations",r.get("observation_date",r.get("date",r.get("DATE"))),freq,sid,r.get("value",r.get(sid)),u,
                          source_notes=r.get("footnotes",""))
    elif s=="silver_bulletin" and p.suffix==".csv":
        for b.line,r in rows(p):
            start,end=iso(r['startdate']),iso(r['enddate'])
            if start>end: b.issue("invalid_interval",r);continue
            sub=r['subgroup'].strip()
            kind=("intensity" if sub in ("Strong","Weak") else "issue" if sub in ("Cost","Economy","Immigration","Trade") else "population_or_pool")
            out=dict(start_date=start,end_date=end,available_at="",president=r['president'],subgroup=sub,subgroup_type=kind,
                     respondent_party="",pollster=r['pollster'].strip(),sponsors=r['sponsors'],population=r['population'].strip(),
                     sample_size=b.numeric(r['samplesize'],"people","samplesize")[0],tracking=r['tracking'],
                     source_created_date=iso(r['createddate']),source_timestamp=r['timestamp'],url=r['url'])
            for k in ('approve','disapprove','net','adjusted_approve','adjusted_disapprove','adjusted_net'):
                out[k+'_fraction']=b.numeric(r[k],"percent",k)[0]
            for k in ('weight','influence'):out['publisher_'+k]=b.numeric(r[k],"publisher_weight",k)[0]
            out['source_record_json']=json.dumps(r,sort_keys=True,separators=(',',':'))
            b.write('survey_observations',out,[r['president'],r['pollster'],start,end,sub,r['population'].strip(),r['samplesize'],r['sponsors']])
    elif s=='cboe':
        for b.line,r in rows(p):
            wide(b,'vix_daily',r,r['DATE'],'day',set(r)-{'DATE'},r['DATE'],drop=['DATE'],units={k:'index_points' for k in r})
    elif s=='french': prepare_french(b,p)
    elif s in ('gpr','epu') and p.suffix in ('.xls','.xlsx'): prepare_index_workbook(b,p)
    elif s in ('epu','infectious_emv') and p.suffix=='.csv':
        for b.line,r in rows(p):
            d=date(int(r['year']),int(r['month']),int(r['day']))
            for k in sorted(set(r)-{'year','month','day'}):b.observation(p.stem,d,'day',k,r[k],'index_points')
    elif s=='michigan':prepare_michigan(b,p)
    elif s=='ucsb':prepare_ucsb(b,p)
    elif s in ('cdc','cdc_nhsn','who','owid','oxcgrt') and p.suffix=='.csv':prepare_health(b,p)
    elif s=='policy_agendas' and p.suffix=='.csv':
        for b.line,r in rows(p):
            b.observation('annual_issue_salience',date(int(r['year']),1,1),'year','share_of_responses',r['percent'],'fraction',
                          group='major_topic_'+r['majortopic'],source_id=r['id'],congress=r['Congress'],
                          basis='publisher_annual_mean_of_normalized_response_shares')
    elif s=='congress':prepare_congress(b,p)
    elif s=='ucdp':prepare_ucdp(b,p)
    elif p.suffix=='.pdf':prepare_pdf(b,p)
    else:
        b.audit('metadata_files',dict(artifact_id=b.a['id'],path=b.a['path'],role='source_metadata_retained'))


def prepare_french(b,p):
    with zipfile.ZipFile(p) as z:
        for member in z.namelist():
            if not member.lower().endswith('.csv'):continue
            text=z.read(member).decode('utf-8-sig');b.audit('source_notes',dict(artifact_id=b.a['id'],member=member,preamble=text.split(',Mkt-RF')[0]))
            for b.line,line in enumerate(text.splitlines(),1):
                r=next(csv.reader([line])); label=r[0].strip() if r else ''
                if not re.fullmatch(r'\d{4}|\d{6}|\d{8}',label) or len(r)!=5:continue
                freq={4:'year',6:'month',8:'day'}[len(label)];d=label if freq=='day' else label+'01' if freq=='month' else label+'0101'
                for name,raw in zip(('Mkt_RF','SMB','HML','RF'),r[1:]):
                    v,u,status=number(raw,'percent',(-99.99,-999))
                    out={**period(d,freq),'series':name,'value':v,'unit':u,'source_value':raw.strip(),'source_unit':'percent_return',
                         'value_status':status,'archive_basis':'latest_reconstructed_history'}
                    out['unit']='fraction_return'
                    b.write('factors_'+freq,out,[d,name])


def prepare_index_workbook(b,p):
    import xlrd
    for sheet,data in workbook(p):
        head=[str(x).strip() for x in data[0]]
        if b.source=='gpr':
            labels={str(r[head.index('var_name')]):str(r[head.index('var_label')]) for r in data[1:] if r[head.index('var_name')]}
            b.audit('source_notes',dict(artifact_id=b.a['id'],sheet=sheet,dictionary=labels))
        else:labels={}
        for i,values in enumerate(data[1:],2):
            b.line=f'{sheet}:{i}';r=dict(zip(head,values))
            if not any(str(v).strip() for v in values):continue
            if b.source=='gpr':
                freq='month' if 'month' in r else 'day'
                if freq=='month':d=xlrd.xldate_as_datetime(float(r['month']),0).date()
                else:d=iso(str(r['DAY']).split('.')[0])
                excluded={'month','DAY','date','var_name','var_label','event'}
                for k,raw in r.items():
                    if k in excluded:continue
                    label=labels.get(k,k)
                    unit=('percent' if 'Percent' in label else 'article_share_source_scale_unverified' if 'Share' in label
                          else 'articles' if k in ('N10','N3H','N10D') else 'index_points')
                    b.observation(p.stem,d,freq,k,raw,unit,definition=label,
                                  geography=k.rsplit('_',1)[-1] if k.startswith(('GPRC_','GPRHC_')) else 'global_news',
                                  publisher_averaging='30_days' if k.endswith('MA30') else '7_days' if k.endswith('MA7') else '')
                if r.get('event'):b.write('event_annotations',{**period(d),'description':r['event'],'precision_note':'publisher_annotation_date'},[iso(d),r['event']])
            else:
                try:d=date(int(float(r['Year'])),int(float(r['Month'])),1)
                except (ValueError,KeyError):
                    b.audit('source_notes',dict(artifact_id=b.a['id'],source_row=b.line,row=r));continue
                for k,raw in r.items():
                    if k not in ('Year','Month'):b.observation(p.stem+'_'+slug(sheet),d,'month',k,raw,'index_points')


def prepare_michigan(b,p):
    suffix=p.stem.split('_')[1];yb=p.stem.startswith('YB')
    measures=['current_economic_index','personal_finances_current','buying_conditions_durables','expected_economic_index',
              'personal_finances_expected','business_conditions_1_year','business_conditions_5_years']
    for sheet,data in workbook(p):
        # Full grid preserves headings, question wording, source notes and unparsed cells.
        b.audit('workbook_grids',dict(artifact_id=b.a['id'],sheet=sheet,rows=data))
        def emit(i,col,d,metric,group='all',unit='index_points',basis='reported_month'):
            b.line=f'{sheet}:R{i+1}C{col+1}'
            canonical={'Democrat':'DEM','Independent':'IND','Republican':'REP','All':'all'}.get(group,group)
            wm=d.year*12+d.month-1-2
            window_start=date(wm//12,wm%12+1,1).isoformat() if basis=='three_month_moving_average' else ''
            b.observation(p.stem,d,'month',metric,data[i][col],unit,canonical,source_group=group,
                          averaging_basis=basis,publisher_window_start=window_start,
                          question_horizon=(('4_years' if d.year<1972 else '5_years') if suffix=='10' else '1_year' if suffix in ('6','8') else ''),
                          transition_note='phone_to_web_Apr_Jul_2024' if iso(d)>='2024-04-01' else '')
        if yb and suffix in ('6','8','10'):
            for col in range(2,len(data[6])):
                d=date(int(data[7][col]),datetime.strptime(data[6][col],'%b').month,1)
                for i in range(8,len(data)):
                    label=str(data[i][0]).strip()
                    if not label or i>39:continue
                    if i in (8,9,10,11,12):emit(i,col,d,slug(label),unit='percent')
                    elif i==13:emit(i,col,d,'sample_size',unit='people')
                    elif i==15:emit(i,col,d,'relative_index')
                    elif i>=22:emit(i,col,d,'relative_index',group=label,basis='three_month_moving_average')
        else:
            for i,row in enumerate(data):
                try:d=date(int(row[1]),datetime.strptime(str(row[0]),'%B').month,1)
                except (ValueError,TypeError,IndexError):continue
                if suffix=='1a':
                    for col,g in ([(4,'all'),(8,'income_under_100000_USD'),(12,'income_over_100000_USD')] if yb else [(2,'all')]):emit(i,col,d,'consumer_sentiment',g)
                    if not yb:emit(i,3,d,'sample_size',unit='people')
                elif suffix=='1b':
                    for col,m in zip(range(2,15,2) if yb else range(2,9),measures):emit(i,col,d,m)
                elif suffix=='5b':
                    cols=[3,4,5,7,8,9,11,12,13] if yb else list(range(2,11))
                    for col,(m,g) in zip(cols,[(m,g) for m in ['consumer_sentiment','current_economic_index','expected_economic_index'] for g in ['DEM','IND','REP']]):
                        emit(i,col,d,m,g,basis='three_month_moving_average' if yb else 'reported_month')
                else:
                    for col,m in enumerate(['better_off','same','worse_off','dk_na','total','relative_index','sample_size'],2):
                        emit(i,col,d,m,unit='percent' if col<=6 else 'people' if col==8 else 'index_points')


def prepare_ucsb(b,p):
    parser=Tables();parser.feed(p.read_text());president=p.stem.replace('-public-approval','')
    for ti,table in enumerate(parser.tables):
        b.audit('html_tables',dict(artifact_id=b.a['id'],table=ti,rows=table))
        if not table or not table[0] or 'start' not in table[0][0].lower():continue
        head=table[0]
        for ri,row in enumerate(table[1:],2):
            b.line=f'table{ti+1}:row{ri}'
            row=row+['']*max(0,len(head)-len(row))
            for col in [2,3,4]+[i for i,h in enumerate(head) if re.search(r'democrat|republican|independent',h,re.I)]:
                if col>=len(row):continue
                party='DEM' if 'democrat' in head[col].lower() else 'REP' if 'republican' in head[col].lower() else 'IND' if 'independent' in head[col].lower() else 'all'
                # Trump I gives independent party fieldwork dates in columns 6/7.
                a,z=(row[6],row[7]) if party!='all' and len(head)>7 and head[6].lower()=='start date' else (row[0],row[1])
                try:start,end=iso(a),iso(z)
                except ValueError:b.issue('invalid_survey_date',dict(header=head[col],row=row));continue
                if not start or not end or start>end:
                    b.issue('invalid_survey_interval',dict(header=head[col],row=row));continue
                metric='approve' if col==2 or party!='all' else 'disapprove' if col==3 else 'unsure'
                v,u,status=b.numeric(row[col],'percent',head[col])
                if v!='' and not 0<=v<=1:b.issue('share_out_of_range',dict(value=v,row=row));status='invalid_range';v=''
                out=dict(start_date=start,end_date=end,available_at='',president=president,respondent_party=party,measure=metric,
                         value=v,unit=u,source_value=row[col],value_status=status,source_column=head[col],sample_size='',population='source_page',
                         source_pollster=row[head.index('Source')] if 'Source' in head else 'see_source_page',
                         comparability_status='party_interpolation_and_methodology_review' if party!='all' else 'source_question_review')
                b.write('approval_observations',out,[president,party,metric,start,end])


def prepare_health(b,p):
    s=b.source;meta={}
    if s=='owid':meta=json.loads((p.parent/'owid_metadata.json').read_text())['fields']
    if s=='cdc_nhsn':meta={x['name']:x for x in json.loads((p.parent/(p.stem+'_metadata.json')).read_text())['columns']}
    for b.line,r in rows(p):
        units={};freq='day';drop=[]
        if s=='who':
            if r.get('Country_code')!='US':b.counts[s+':non_US_rows_excluded']+=1;continue
            d=r['Date_reported'];freq='report_date';num={'New_cases','Cumulative_cases','New_deaths','Cumulative_deaths'};key=[d,r['Country_code']]
            units={k:'reported_count' for k in num}
        elif s=='owid':
            if r['country']!='United States':b.counts[s+':non_US_rows_excluded']+=1;continue
            d=r['date'];num=set(r)-{'country','date','code','continent'};key=[d,r['country']]
            units={k:('percent' if meta.get(k,{}).get('unit') in ('%','% of population ages 20 to 79') else meta.get(k,{}).get('unit') or 'source_numeric') for k in num}
        elif s=='oxcgrt':
            if r['CountryCode']!='USA':b.counts[s+':non_US_rows_excluded']+=1;continue
            d=r['Date'];textcols={'CountryName','CountryCode','RegionName','RegionCode','CityName','CityCode','Jurisdiction','Date'}
            textcols |= {k for k in r if k=='MajorityVaccinated' or k.startswith(('V2B_','V2C_'))}
            num=set(r)-textcols;key=[d,*[r.get(k,'') for k in ['RegionCode','CityCode','Jurisdiction']]]
            units={k:'ordinal_policy_code' if re.match(r'[CEHV]\d',k) else 'index_points' if 'Index' in k else 'source_numeric' for k in num}
            for k in num:
                if k.startswith(('E3_','E4_','H4_','H5_')):units[k]='USD_reported_fiscal_spending'
                elif k.endswith('_Flag'):units[k]='binary_scope_flag'
                elif k.startswith('V1_'):units[k]='categorical_vaccine_prioritization_code'
                elif k=='PopulationVaccinated':units[k]='percent'
                elif k in ('ConfirmedCases','ConfirmedDeaths'):units[k]='reported_cumulative_count'
        elif s=='cdc_nhsn':
            d=r['Week Ending Date'];freq='week';num={k for k in r if meta.get(k,{}).get('dataTypeName')=='number'};key=[d,r['Geographic aggregation']]
            for k in num:units[k]='percent' if k.lower().startswith('percent') else 'source_reported_rate' if 'rate' in k.lower() else 'source_reported_number'
        else:
            d=r['Week Ending Date'];freq='week';num={'Observed Number','Upper Bound Threshold','Average Expected Count','Excess Estimate','Total Excess Estimate','Percent Excess Estimate'}
            key=[d,r['State'],r['Type'],r['Outcome']]
            units={k:'deaths' for k in num};units['Percent Excess Estimate']='percent';units['Total Excess Estimate']='publisher_multiweek_total_not_weekly_flow'
        wide(b,p.stem,r,d,freq,num,key,drop,units)


def prepare_congress(b,p):
    for i,person in enumerate(json.loads(p.read_text()),1):
        b.line=i;pid=person['id']['bioguide']
        b.write('people',dict(person_id=pid,name_json=json.dumps(person['name'],sort_keys=True),identifiers_json=json.dumps(person['id'],sort_keys=True)),pid)
        for j,t in enumerate(person['terms'],1):
            b.line=f'person{i}:term{j}'
            start,end=iso(t['start']),iso(t['end'])
            if start>end:b.issue('invalid_term_interval',t);continue
            b.write('office_terms',dict(person_id=pid,start_date=start,end_date=end,available_at='',end_boundary='source_unspecified',
                                       chamber=t['type'],state=t['state'],party=t.get('party',''),seat_class=t.get('class',''),
                                       source_term_json=json.dumps(t,sort_keys=True)),[pid,t['type'],t['state'],start,end])


def prepare_ucdp(b,p):
    def emit(reader):
        for b.line,r in enumerate(reader,2):
            a,z=iso(r['date_start']),iso(r['date_end'])
            if not a or not z or a>z:b.issue('invalid_event_interval',r);continue
            out={k:(v.strip() if isinstance(v,str) else v) for k,v in r.items()}
            out['date_start']=a;out['date_end']=z
            for k in ['deaths_a','deaths_b','deaths_civilians','deaths_unknown','best','high','low','latitude','longitude']:
                out[k]=b.numeric(r[k],'degrees' if k in ('latitude','longitude') else 'deaths',k)[0]
            out['available_at']='';out['archive_kind']='annual_finalized' if p.suffix=='.zip' else 'candidate_provisional'
            b.write('event_versions',out,r['id'])
    if p.suffix=='.zip':
        with zipfile.ZipFile(p) as z:
            for m in z.namelist():
                if m.lower().endswith('.csv'):
                    with z.open(m) as f:emit(csv.DictReader(io.TextIOWrapper(f,encoding='utf-8-sig')))
    else:
        with p.open(encoding='utf-8-sig',newline='') as f:emit(csv.DictReader(f))


def prepare_pdf(b,p):
    from pypdf import PdfReader
    reader=PdfReader(p);first=reader.pages[0].extract_text() or ''
    m=re.search(r'([A-Z][a-z]+) (\d{1,2})\s*-\s*(\d{1,2}), (20\d\d)\s*-\s*(\d+) U\.S\. ([^\n]+)',first)
    dates={}
    if m:
        mo,a,z,yr,n,pop=m.groups();month=datetime.strptime(mo,'%B').month
        dates=dict(start_date=date(int(yr),month,int(a)).isoformat(),end_date=date(int(yr),month,int(z)).isoformat(),sample_size=int(n),population=pop.strip())
    b.audit('pdf_documents',dict(artifact_id=b.a['id'],pages=len(reader.pages),**dates,
                               structured_status='selected_simple_toplines_only' if 'Sample' in first else 'page_text_for_review'))
    for page,pg in enumerate(reader.pages,1):
        b.line=page;text=pg.extract_text() or ''
        b.audit('document_pages',dict(artifact_id=b.a['id'],page=page,text=text))
        # A deliberately narrow extraction: numbered, scalar dot-leader answer blocks.
        # Cross-tabs/matrices stay searchable text; never infer their subgroup columns.
        if b.source!='yougov' or 'Sample' not in first or not dates:continue
        for block in re.split(r'(?m)^(?=\d+[A-Z]?\. )',text):
            qm=re.match(r'(\d+[A-Z]?)\. ([\s\S]+)',block)
            if not qm:continue
            qid,body=qm.groups();lines=body.splitlines();answers=[];q=[];started=False
            for line in lines:
                am=re.fullmatch(r'(.+?)\s*(?:\.\s*){3,}\s*(\d+(?:\.\d+)?)%',line.strip())
                if am:answers.append(am.groups());started=True
                elif not started:q.append(line)
            if not answers:continue
            question=' '.join(q).strip()
            # Extract only complete, single-page scalar questions; retain filters in wording.
            for answer,value in answers:
                out={**dates,'available_at':'','page':page,'question_number':qid,'question':question,
                     'answer':answer.strip(),'value':float(value)/100,'unit':'fraction','source_value':value+'%',
                     'denominator_status':'question_filters_require_review','sample_size_scope':'whole_report_not_question_or_subgroup',
                     'answer_role':'publisher_total' if answer.lower().startswith('total ') else 'response_option',
                     'extraction_status':'scalar_topline_review'}
                b.write('topline_candidates',out,[b.a['id'],page,qid,answer])


def pin_inputs(config, refresh=False):
    if config.exists() and not refresh:return json.loads(config.read_text())
    bundles=[];unique={}
    for pointer in sorted(BASE.glob('*/*/latest.json')):
        info=json.loads(pointer.read_text());p=pointer.parent/info['snapshot']
        manifest=json.loads((p/'manifest.json').read_text());source=pointer.parent.parent.name
        bundles.append(dict(path=str(p.relative_to(LAB)),manifest_sha256=sha_file(p/'manifest.json')))
        for name,item in sorted(manifest['files'].items()):
            if not name.startswith('raw/'):continue
            key=(source,Path(name).name,item['sha256'])
            origin=dict(path=str((p/name).relative_to(LAB)),url=item.get('url',''),retrieved_at=item.get('retrieved_at',manifest.get('retrieved_at','')))
            if key not in unique:
                unique[key]=dict(id=stable(key)[:24],source=source,path=origin['path'],sha256=item['sha256'],origins=[])
            unique[key]['origins'].append(origin)
    if not bundles:raise ValueError('No downloaded feature bundles')
    result=dict(schema_version=VERSION,bundles=bundles,artifacts=list(unique.values()))
    config.parent.mkdir(parents=True,exist_ok=True);atomic_json(config,result);return result


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--refresh-inputs',action='store_true');parser.add_argument('--config',type=Path,default=CONFIG)
    parser.add_argument('--output',type=Path,default=OUTPUT);args=parser.parse_args()
    inputs=pin_inputs(args.config,args.refresh_inputs)
    for entry in inputs['bundles']:
        if sha_file(LAB/entry['path']/'manifest.json')!=entry['manifest_sha256']:raise ValueError('Pinned manifest changed')
    with Snapshot(args.output,'source_separated_prepared_features') as snap:
        snap.write_json('inputs.json',inputs)
        snap.add('recipe/prepare_features.py',Path(__file__).read_bytes())
        snap.add('recipe/data_utils.py',(Path(__file__).parent/'data_utils.py').read_bytes())
        from importlib.metadata import version
        snap.write_json('recipe/dependencies.json',{k:version(k) for k in ('xlrd','openpyxl','pypdf')})
        snap.write_json('preparation.json',dict(version=VERSION,code_sha256=sha_file(Path(__file__)),
                                               aligned_to_polls=False,forecast_ready=False,release_dates_imputed=False))
        b=Builder(snap)
        for a in inputs['artifacts']:b.run(a)
        groups=b.finish()
        summary=dict(sources=len({a['source'] for a in inputs['artifacts']}),input_bundles=len(inputs['bundles']),
                     unique_raw_artifacts=len(inputs['artifacts']),tables=len(b.catalog),
                     rows=sum(c['rows'] for c in b.catalog.values()),overlapping_variant_groups=groups,
                     counters=dict(sorted(b.counts.items())),aligned_to_polls=False)
        snap.write_json('catalog.json',b.catalog);snap.write_json('summary.json',summary)
        for p in sorted(snap.stage.rglob('*')):
            if p.is_file():snap.files[str(p.relative_to(snap.stage))]=dict(sha256=sha_file(p),bytes=p.stat().st_size)
        snap.finish({k:v for k,v in summary.items() if k!='counters'})

if __name__=='__main__':main()
